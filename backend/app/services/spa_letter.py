"""Geração da minuta do Ofício à Secretaria de Prêmios e Apostas (SPA/MF).

Gera um documento .docx editável (minuta) comunicando à SPA a relação de agentes operadores
inadimplentes em determinado mês de competência, conforme o modelo institucional da confederação.
"""
import os
import uuid
import logging
from datetime import date

logger = logging.getLogger(__name__)

MESES = [
    "", "janeiro", "fevereiro", "março", "abril", "maio", "junho",
    "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
]

UPLOAD_DIR = "/app/uploads/oficios"


def _fmt_mes_ano(d: date) -> str:
    return f"{MESES[d.month]}/{d.year}"


def _fmt_data_extenso(d: date) -> str:
    return f"{d.day:02d} de {MESES[d.month]} de {d.year}"


def build_spa_letter_text(
    confederation,
    reference_month: date,
    inadimplentes: list,
    endr_count: int = 0,
    city: str = "Rio de Janeiro",
    letter_date: date = None,
    first_notif_date: str = "",
    second_notif_date: str = "",
    spa_list_date: str = "",
    endr_list_date: str = "",
) -> str:
    """Monta o texto da minuta do ofício (também usado como pré-visualização)."""
    letter_date = letter_date or date.today()
    mes_ano = _fmt_mes_ano(reference_month)
    conf_nome = confederation.name
    conf_sigla = confederation.acronym

    linhas = []
    linhas.append(f"{city}, {_fmt_data_extenso(letter_date)}")
    linhas.append("")
    linhas.append("À Secretaria de Prêmios e Apostas – SPA")
    linhas.append("Ministério da Fazenda")
    linhas.append("Esplanada dos Ministérios – Bloco P – 2º andar – Brasília/DF")
    linhas.append("")
    linhas.append(
        "Assunto: Comunicação sobre inadimplência de agentes operadores em relação ao "
        "art. 30, §1º-A, III, “a”, da Lei nº 13.756/2018."
    )
    linhas.append("")
    linhas.append(
        f"A {conf_nome} ({conf_sigla}), na condição de entidade beneficiária dos repasses previstos "
        f"no art. 30, §1º-A, III, “a”, da Lei nº 13.756/2018, dirige-se a este órgão regulador para "
        f"tratar dos agentes operadores que não realizaram os repasses obrigatórios durante o mês de {mes_ano}."
    )
    linhas.append("")
    linhas.append(
        "Conforme o regramento introduzido no início de 2025, foi estabelecido procedimento padronizado "
        "para a transferência das parcelas devidas às entidades beneficiárias. A Portaria SPA/MF nº 41/2025 "
        "determinou que os agentes operadores realizassem o repasse direto aos entes contemplados pela "
        "legislação, em frequência mensal, a partir de 1º de janeiro de 2025. Esse fluxo passou a vigorar de "
        "forma efetiva em 31 de janeiro de 2025, com prazos de vencimento fixados para o dia 10 do mês "
        "subsequente à apuração, conforme dispõem os arts. 3º e 6º da Instrução Normativa SPA/MF nº 9/2025."
    )
    linhas.append("")
    notif_txt = (
        f"foram enviados 02 (dois) e-mails a cada agente operador inadimplente reforçando a necessidade de "
        f"repasse dos valores"
    )
    if first_notif_date or second_notif_date:
        notif_txt += f", o primeiro no dia {first_notif_date or '__/__'} e o segundo no dia {second_notif_date or '__/__'}"
    linhas.append(
        f"Ressalta-se que a {conf_sigla} adotou as providências administrativas e estruturais necessárias "
        f"para o recebimento dos repasses, bem como empreendeu esforços de contato com os agentes operadores, "
        f"com o objetivo de viabilizar o fiel cumprimento da legislação aplicável. Operacionalmente, {notif_txt}. "
        f"Entretanto, verifica-se que parte dos operadores ainda não vem observando as obrigações regulamentares, "
        f"deixando de efetuar os repasses devidos dentro dos marcos temporais estabelecidos."
    )
    linhas.append("")
    linhas.append(
        f"Para identificação dos operadores inadimplentes, a {conf_sigla} considerou: "
        f"(i) a relação de operadores autorizados divulgada pela SPA"
        + (f" (publicada no dia {spa_list_date})" if spa_list_date else "")
        + "; (ii) a exclusão dos operadores associados ao Escritório Nacional de Rateio - ENDR"
        + (f" (conforme lista apresentada em {endr_list_date})" if endr_list_date else "")
        + f"; e (iii) a exclusão daqueles que efetuaram repasse no mês de {mes_ano}. "
        f"Ao final desse procedimento, obteve-se a seguinte relação de agentes operadores inadimplentes:"
    )
    linhas.append("")
    linhas.append("[TABELA: Autorização | CNPJ | Razão social]")
    for it in inadimplentes:
        linhas.append(
            f"{it.get('autorizacao') or '—'}\t{it.get('cnpj') or '—'}\t{it.get('razao_social') or '—'}"
        )
    linhas.append("")
    if endr_count:
        linhas.append(
            f"No que se refere aos agentes operadores associados ao ENDR, também não foram realizados "
            f"repasses durante o mês de {mes_ano}."
        )
        linhas.append("")
    linhas.append(
        "O marco regulatório atribui responsabilidade direta ao agente operador pelo repasse dos valores "
        "legais, prevendo sanções de natureza civil, administrativa e criminal pelo descumprimento "
        "(Portaria SPA/MF nº 1.240/2024). Ainda, as Portarias SPA/MF nº 1.225/2024 e nº 1.233/2024 conferem "
        "à Administração instrumentos sancionatórios progressivos, incluindo advertência, multa, suspensão e "
        "eventual cassação da autorização."
    )
    linhas.append("")
    linhas.append(
        f"Diante desse cenário, a {conf_sigla} solicita adoção das medidas regulatórias cabíveis para garantir "
        f"o cumprimento dos repasses obrigatórios e esclarecimento formal sobre eventuais alegações de operação "
        f"inativa ou GGR negativo por parte dos operadores inadimplentes."
    )
    linhas.append("")
    linhas.append(
        "Reiteramos a confiança na atuação dessa Secretaria para a adoção das medidas cabíveis e permanecemos "
        "à disposição para prestar informações adicionais e acompanhar os desdobramentos necessários."
    )
    linhas.append("")
    linhas.append("Atenciosamente,")
    linhas.append("")
    linhas.append(confederation.president_name or "[Nome do Presidente]")
    linhas.append(f"Presidente – {conf_sigla}")
    return "\n".join(linhas)


def generate_spa_letter_docx(
    confederation,
    reference_month: date,
    inadimplentes: list,
    endr_count: int = 0,
    city: str = "Rio de Janeiro",
    letter_date: date = None,
    first_notif_date: str = "",
    second_notif_date: str = "",
    spa_list_date: str = "",
    endr_list_date: str = "",
) -> dict:
    """Gera o arquivo .docx da minuta e retorna {file_path, file_name, text}."""
    from docx import Document as Docx
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    os.makedirs(UPLOAD_DIR, exist_ok=True)
    letter_date = letter_date or date.today()
    text = build_spa_letter_text(
        confederation, reference_month, inadimplentes, endr_count, city, letter_date,
        first_notif_date, second_notif_date, spa_list_date, endr_list_date,
    )

    doc = Docx()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    mes_ano = _fmt_mes_ano(reference_month)
    conf_sigla = confederation.acronym

    p = doc.add_paragraph(f"{city}, {_fmt_data_extenso(letter_date)}")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph("À Secretaria de Prêmios e Apostas – SPA")
    doc.add_paragraph("Ministério da Fazenda")
    doc.add_paragraph("Esplanada dos Ministérios – Bloco P – 2º andar – Brasília/DF")
    doc.add_paragraph("")

    pa = doc.add_paragraph()
    run = pa.add_run(
        "Assunto: Comunicação sobre inadimplência de agentes operadores em relação ao "
        "art. 30, §1º-A, III, “a”, da Lei nº 13.756/2018."
    )
    run.bold = True

    # Corpo (parágrafos justificados) — reaproveita o texto, exceto cabeçalho/tabela/assinatura.
    for bloco in text.split("\n\n"):
        bloco = bloco.strip()
        if not bloco:
            continue
        if bloco.startswith(("Rio de Janeiro", city)) or bloco.startswith("À Secretaria") \
           or bloco.startswith("Ministério") or bloco.startswith("Esplanada") \
           or bloco.startswith("Assunto:") or bloco.startswith("[TABELA"):
            continue
        if bloco.startswith("Atenciosamente"):
            break
        par = doc.add_paragraph(bloco)
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Tabela de inadimplentes
    if inadimplentes:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Autorização"
        hdr[1].text = "CNPJ"
        hdr[2].text = "Razão social"
        for c in hdr:
            for par in c.paragraphs:
                for r in par.runs:
                    r.bold = True
        for it in inadimplentes:
            row = table.add_row().cells
            row[0].text = str(it.get("autorizacao") or "—")
            row[1].text = str(it.get("cnpj") or "—")
            row[2].text = str(it.get("razao_social") or "—")

    doc.add_paragraph("")
    if endr_count:
        par = doc.add_paragraph(
            f"No que se refere aos agentes operadores associados ao ENDR, também não foram realizados "
            f"repasses durante o mês de {mes_ano}."
        )
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for bloco in [
        "O marco regulatório atribui responsabilidade direta ao agente operador pelo repasse dos valores "
        "legais, prevendo sanções de natureza civil, administrativa e criminal pelo descumprimento "
        "(Portaria SPA/MF nº 1.240/2024). Ainda, as Portarias SPA/MF nº 1.225/2024 e nº 1.233/2024 conferem "
        "à Administração instrumentos sancionatórios progressivos, incluindo advertência, multa, suspensão e "
        "eventual cassação da autorização.",
        f"Diante desse cenário, a {conf_sigla} solicita adoção das medidas regulatórias cabíveis para garantir "
        f"o cumprimento dos repasses obrigatórios e esclarecimento formal sobre eventuais alegações de operação "
        f"inativa ou GGR negativo por parte dos operadores inadimplentes.",
        "Reiteramos a confiança na atuação dessa Secretaria para a adoção das medidas cabíveis e permanecemos "
        "à disposição para prestar informações adicionais e acompanhar os desdobramentos necessários.",
    ]:
        par = doc.add_paragraph(bloco)
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_paragraph("")
    doc.add_paragraph("Atenciosamente,")
    doc.add_paragraph("")
    doc.add_paragraph(confederation.president_name or "[Nome do Presidente]")
    doc.add_paragraph(f"Presidente – {conf_sigla}")

    file_name = f"Oficio_SPA_{conf_sigla}_{reference_month.strftime('%Y_%m')}.docx"
    file_path = os.path.join(UPLOAD_DIR, f"{uuid.uuid4().hex}_{file_name}")
    doc.save(file_path)

    return {"file_path": file_path, "file_name": file_name, "text": text}
